from __future__ import annotations

import json
import hashlib
import logging
from pathlib import Path

from pptx import Presentation

from ppt_agent.coordinator.phase_state import load_artifact, write_artifact
from ppt_agent.models.artifacts import JobWorkspace
from ppt_agent.models.validation import validation_report

logger = logging.getLogger(__name__)


def _strict_assembly_invariants(
    final_pptx: Path,
    template_zones: dict,
    slide_contents: dict,
) -> dict:
    """Verify that strict assembly changed text/image payloads only.

    Geometry and formatting are compared with fingerprints captured during
    template ingestion.  Slides in the output are aligned to the selected
    template pages recorded by the mapper.
    """
    from ppt_agent.templates.ingest import _extract_shape_formatting, _iter_shapes_recursive

    result = {"applicable": False, "passed": True, "checked_zones": 0, "issues": []}
    if slide_contents.get("assembly_policy", {}).get("mode") != "text_replace_only":
        return result
    if not final_pptx.exists():
        return result

    template_by_index = {
        int(slide.get("index", index)): slide
        for index, slide in enumerate(template_zones.get("slides", []))
    }
    if not any(
        zone.get("geometry_fingerprint") or zone.get("formatting_fingerprint")
        for slide in template_by_index.values()
        for zone in slide.get("all_zones", [])
    ):
        return result

    result["applicable"] = True
    presentation = Presentation(str(final_pptx))
    for output_index, mapped_slide in enumerate(slide_contents.get("slides", [])):
        if output_index >= len(presentation.slides):
            result["issues"].append({
                "slide_index": output_index,
                "kind": "missing_output_slide",
            })
            continue
        template_index = int(mapped_slide.get("template_slide_index", output_index))
        template_slide = template_by_index.get(template_index, {})
        shape_items = list(_iter_shapes_recursive(
            presentation.slides[output_index].shapes, template_index
        ))
        shapes_by_path = {item["shape_path"]: item for item in shape_items}
        shapes_by_id = {
            int(item["shape"].shape_id): item for item in shape_items
        }
        for zone in template_slide.get("all_zones", []):
            expected_geometry = zone.get("geometry_fingerprint")
            # Font/paragraph formatting only applies to editable text shapes.
            # Non-text zones store an empty-dict fingerprint during ingestion;
            # comparing that to a text-format default dict creates false drift.
            expected_formatting = (
                zone.get("formatting_fingerprint") if zone.get("editable") else None
            )
            if not expected_geometry and not expected_formatting:
                continue
            result["checked_zones"] += 1
            native_shape_id = zone.get("native_shape_id")
            item = shapes_by_path.get(zone.get("shape_path") or zone.get("zone_id"))
            if item is None and native_shape_id is not None:
                item = shapes_by_id.get(int(native_shape_id))
            if item is None:
                result["issues"].append({
                    "slide_index": output_index,
                    "template_slide_index": template_index,
                    "zone_id": zone.get("zone_id", ""),
                    "kind": "missing_shape",
                })
                continue
            shape = item["shape"]

            geometry_payload = ":".join(
                str(round(value, 3)) for value in item["displayed_geometry"]
            )
            actual_geometry = hashlib.sha256(
                geometry_payload.encode("utf-8")
            ).hexdigest()[:16]
            if expected_geometry and actual_geometry != expected_geometry:
                result["issues"].append({
                    "slide_index": output_index,
                    "template_slide_index": template_index,
                    "zone_id": zone.get("zone_id", ""),
                    "kind": "geometry_drift",
                    "expected": expected_geometry,
                    "actual": actual_geometry,
                })

            formatting_payload = json.dumps(
                _extract_shape_formatting(shape), sort_keys=True, ensure_ascii=False
            )
            actual_formatting = hashlib.sha256(
                formatting_payload.encode("utf-8")
            ).hexdigest()[:16]
            if expected_formatting and actual_formatting != expected_formatting:
                result["issues"].append({
                    "slide_index": output_index,
                    "template_slide_index": template_index,
                    "zone_id": zone.get("zone_id", ""),
                    "kind": "formatting_drift",
                    "expected": expected_formatting,
                    "actual": actual_formatting,
                })

    result["passed"] = not result["issues"]
    return result

# ── Visual audit integration ──
try:
    from ppt_agent.vision.pptx_audit import audit_pptx as visual_audit_pptx
except ImportError:
    visual_audit_pptx = None
    logger.info("Vision module not available — visual audit skipped")

_FRESH_EYES_SYSTEM = (
    "You are an independent PPT quality reviewer. You have NOT seen any prior "
    "conversation or phase context — you are reviewing the final output with "
    "fresh eyes.\n\n"
    "Evaluate the presentation on three dimensions:\n"
    "1. **Content completeness**: Does slide_contents.json cover all expected "
    "slides, sections, and key points? Are any slides empty or missing?\n"
    "2. **Visual consistency**: Are fallback flags resolved? Are image "
    "references present and plausible? Does the image_generation_report "
    "(if provided) indicate failures?\n"
    "3. **Text accuracy**: Is the text in slide_contents.json coherent, free "
    "of obvious errors, and does it read like professional presentation copy?\n\n"
    "Return your findings as a JSON object with keys:\n"
    "  - content_completeness: {score: 0-100, issues: [...]}\n"
    "  - visual_consistency:   {score: 0-100, issues: [...]}\n"
    "  - text_accuracy:        {score: 0-100, issues: [...]}\n"
    "  - overall_score: 0-100\n"
    "  - summary: string\n"
)


def _fresh_eyes_validation(workspace: JobWorkspace, llm_client) -> dict:
    """Run an LLM-based validation with isolated (fresh-eyes) context.

    The LLM receives ONLY the final artifacts — no accumulated conversation
    history from prior phases.
    """
    slide_contents = load_artifact(workspace, "slide_contents")

    # Build a minimal context containing only the artifacts needed
    context: dict = {
        "slide_contents": slide_contents,
        "final_pptx_exists": (workspace.root / "final.pptx").exists(),
    }

    image_report_path = workspace.artifact_path("image_generation_report")
    if image_report_path.exists():
        context["image_generation_report"] = load_artifact(workspace, "image_generation_report")

    # Deterministic fallback in case the LLM call fails
    deterministic_warnings: list[str] = []
    if context.get("image_generation_report"):
        deterministic_warnings.extend(context["image_generation_report"].get("warnings", []))
    if not context["final_pptx_exists"]:
        deterministic_warnings.append("final.pptx is missing")

    fallback = {
        "content_completeness": {"score": 70, "issues": []},
        "visual_consistency": {"score": 70, "issues": deterministic_warnings},
        "text_accuracy": {"score": 70, "issues": []},
        "overall_score": 70,
        "summary": "Deterministic fallback — LLM fresh-eyes validation was unavailable.",
    }

    image_report = context.get("image_generation_report") or {}
    generated_background_mode = (
        slide_contents.get("assembly_policy", {}).get("mode")
        == "generated_background_replace"
    )
    if generated_background_mode and workspace.artifact_path("outline").exists():
        # Full-page background assembly deliberately builds its editable text
        # layer from the outline rather than fragmented source-template zones.
        # Review the content that is actually rendered.
        outline = load_artifact(workspace, "outline")
        compact_slides = [
            {
                "slide_index": slide.get("slide_index", index),
                "title": slide.get("title", ""),
                "content": slide.get("bullets", []),
                "fallback_flags": [],
            }
            for index, slide in enumerate(outline.get("slides", []))
        ]
    else:
        compact_slides = _compact_mapped_slides_for_review(
            slide_contents.get("slides", [])
        )
    compact_contents = {
        "slide_count": len(compact_slides),
        "review_status": slide_contents.get("review_status"),
        "mapping_mode": slide_contents.get("mapping_mode"),
        "generation_stats": slide_contents.get("generation_stats"),
        "assembly_text_source": (
            "outline" if generated_background_mode else "mapped_template_zones"
        ),
        "slides": compact_slides,
    }

    prompt = (
        "Review the following PPT artifacts and provide your quality assessment.\n\n"
        "The following is a compact, complete projection of slide_contents.json; "
        "it is not truncated. It contains deduplicated semantic copy that is "
        "actually mapped into the inherited template. Decorative glyphs, repeated "
        "shadow/outline layers, and preserved template labels are intentionally "
        "excluded; do not treat their absence as missing content.\n"
        f"slide_contents.json:\n```json\n{json.dumps(compact_contents, ensure_ascii=False, indent=2)}\n```\n\n"
        f"final.pptx exists: {context['final_pptx_exists']}\n"
    )
    if "image_generation_report" in context:
        report_text = json.dumps(context["image_generation_report"], ensure_ascii=False, indent=2)[:4000]
        prompt += f"\nimage_generation_report.json:\n```json\n{report_text}\n```\n"

    logger.info("Running fresh-eyes LLM validation (no prior phase context)")

    from ppt_agent.llm.schemas import VERIFICATION_SCHEMA

    result = llm_client.generate_json(
        prompt=prompt,
        context={},  # Intentionally empty — no prior phase context
        system=_FRESH_EYES_SYSTEM,
        phase="verification",
        fallback=fallback,
        schema=VERIFICATION_SCHEMA,
    )

    # Ensure expected keys are present
    result.setdefault("content_completeness", {"score": 70, "issues": []})
    result.setdefault("visual_consistency", {"score": 70, "issues": []})
    result.setdefault("text_accuracy", {"score": 70, "issues": []})
    result.setdefault("overall_score", 70)
    result.setdefault("summary", "")

    return result


def _compact_mapped_slides_for_review(slides: list[dict]) -> list[dict]:
    """Project mapped template zones into semantic copy for fresh-eyes review.

    Authored templates often represent one visible label with several overlapping
    shapes (fill, outline and shadow), and may split a decorative word into
    one-character shapes. Those implementation details are useful for assembly but
    misleading to a content reviewer, so the projection keeps only active,
    meaningful replacement copy and deduplicates layered text.
    """

    compact_slides: list[dict] = []
    excluded_types = {"decorative", "footer", "page_number", "noise"}

    for index, slide in enumerate(slides):
        seen: set[str] = set()
        copy: list[dict] = []

        for zone in slide.get("zones", []):
            content = str(zone.get("content") or "").strip()
            if not content or zone.get("action") != "replace_text":
                continue
            zone_type = str(zone.get("type") or "body").lower()
            semantic_role = str(zone.get("semantic_role") or "").lower()
            eligibility = str(zone.get("content_eligibility") or "").lower()
            if (
                zone_type in excluded_types
                or semantic_role in {"template_noise", "decorative"}
                or eligibility in {"decorative", "template_noise"}
            ):
                continue

            normalized = "".join(content.split()).casefold()
            if len(normalized) <= 1 or normalized in seen:
                continue
            seen.add(normalized)
            copy.append(
                {
                    "role": semantic_role or zone_type,
                    "text": content,
                }
            )

        title = next(
            (
                item["text"]
                for item in copy
                if item["role"] in {"slide_title", "title"}
            ),
            copy[0]["text"] if copy else "",
        )
        compact_slides.append(
            {
                "slide_index": slide.get("slide_index", index),
                "title": title,
                "key_copy": copy,
                "fallback_flags": slide.get("fallback_flags", []),
            }
        )

    return compact_slides


def run(workspace: JobWorkspace, force: bool = False, llm_client=None) -> Path:
    output = workspace.artifact_path("validation_report")
    if output.exists() and not force:
        dependencies = [
            workspace.root / "final.pptx",
            workspace.artifact_path("image_generation_report"),
        ]
        if all(
            not dependency.exists()
            or output.stat().st_mtime >= dependency.stat().st_mtime
            for dependency in dependencies
        ):
            return output

    slide_contents = load_artifact(workspace, "slide_contents")

    # Collect deterministic warnings (always run)
    warnings: list[str] = []
    image_report_path = workspace.artifact_path("image_generation_report")
    image_report = None
    if image_report_path.exists():
        image_report = load_artifact(workspace, "image_generation_report")
        warnings.extend(image_report.get("warnings", []))
    if not (workspace.root / "final.pptx").exists():
        warnings.append("final.pptx is missing")

    report = validation_report(workspace.root.name, slide_contents, warnings)

    # Strict text-only assembly must be mechanically provable: no shape
    # movement/resizing and no typography drift relative to ingested zones.
    final_pptx = workspace.root / "final.pptx"
    template_zones_path = workspace.artifact_path("template_zones")
    generated_background_mode = (
        slide_contents.get("assembly_policy", {}).get("mode")
        == "generated_background_replace"
    )
    if generated_background_mode:
        # The mapped template zones are not the final editable text layer in
        # this mode. The assembler creates fresh PPT text boxes over each
        # background, so template-zone editability flags are not applicable.
        for slide_result in report.get("slide_results", []):
            slide_result["issues"] = [
                issue
                for issue in slide_result.get("issues", [])
                if issue != "One or more text zones are not editable."
            ]
            slide_result.setdefault("checks", {})["text_editable"] = True
            slide_result["status"] = (
                "passed" if not slide_result["issues"] else "warning"
            )
        if (
            not report.get("manual_review_items")
            and all(
                item.get("status") == "passed"
                for item in report.get("slide_results", [])
            )
        ):
            report["status"] = "passed"

    if template_zones_path.exists() and not generated_background_mode:
        invariant_report = _strict_assembly_invariants(
            final_pptx,
            load_artifact(workspace, "template_zones"),
            slide_contents,
        )
        report["strict_assembly_invariants"] = invariant_report
        if invariant_report["applicable"] and not invariant_report["passed"]:
            report["status"] = "failed"
            message = (
                "Strict assembly invariant failure: "
                f"{len(invariant_report['issues'])} geometry/format issue(s)."
            )
            report["manual_review_items"].append(message)
    elif generated_background_mode:
        report["strict_assembly_invariants"] = {
            "applicable": False,
            "passed": True,
            "checked_zones": 0,
            "issues": [],
            "reason": "Full-page generated backgrounds intentionally replace template geometry.",
        }

    generation_required = False
    design_plan_path = workspace.artifact_path("slide_design_plan")
    if design_plan_path.exists():
        design_plan = load_artifact(workspace, "slide_design_plan")
        generation_required = any(
            str(slide.get("visual_strategy") or "").lower()
            in {"generated_image", "generated_background"}
            for slide in design_plan.get("slides", [])
        )
    if generation_required and image_report and image_report.get("total_slides", 0):
        generated = int(image_report.get("generated", 0) or 0)
        if generated == 0:
            report["status"] = "failed"
            report["manual_review_items"].append(
                "Visual generation produced no usable slide backgrounds."
            )

    # ── Visual audit (PPTX structural check, no image rendering needed) ──
    if final_pptx.exists() and visual_audit_pptx is not None:
        try:
            visual_report = visual_audit_pptx(final_pptx)
            report["visual_audit"] = visual_report.to_dict()
            # Promote visual issues to main warnings
            for rec in visual_report.recommendations:
                warnings.append(f"[视觉检查] {rec}")
            logger.info(
                "Visual audit: score=%.1f, font_issues=%d, contrast_issues=%d, overflow_issues=%d",
                visual_report.aggregate_score,
                visual_report.total_font_issues,
                visual_report.total_contrast_issues,
                visual_report.total_overflow_issues,
            )
        except Exception as exc:
            logger.warning("Visual audit failed: %s", exc)
    elif visual_audit_pptx is None:
        logger.debug("Visual audit skipped — vision module not available")

    # If an LLM is available, augment with fresh-eyes validation
    if llm_client is not None:
        try:
            fresh_eyes = _fresh_eyes_validation(workspace, llm_client)
            report["fresh_eyes_validation"] = fresh_eyes
            if int(fresh_eyes.get("overall_score", 0) or 0) < 75:
                report["status"] = "failed"
                report["manual_review_items"].append(
                    "Fresh-eyes quality score is below the release threshold of 75."
                )
            logger.info(
                "Fresh-eyes validation complete: overall_score=%s",
                fresh_eyes.get("overall_score"),
            )
        except Exception as exc:
            logger.warning("Fresh-eyes LLM validation failed, using deterministic only: %s", exc)
    else:
        logger.info("No LLM client available — using deterministic verification only")

    return write_artifact(workspace, "validation_report", report)
