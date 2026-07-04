from __future__ import annotations

import logging
from pathlib import Path

from pptx import Presentation

logger = logging.getLogger(__name__)


def inspect_pptx(path: Path, expected_slide_count: int | None = None) -> dict:
    """Inspect a PPTX file and return a verification report.

    Parameters
    ----------
    path:
        Path to the ``.pptx`` file.
    expected_slide_count:
        If provided, a warning is emitted when the actual slide count differs.

    Returns
    -------
    dict
        Keys: ``slide_count``, ``total_shapes``, ``per_slide_info``,
        ``warnings``, ``file_size_bytes``, ``exists``.
    """
    result: dict = {
        "exists": path.exists(),
        "file_size_bytes": 0,
        "slide_count": 0,
        "total_shapes": 0,
        "per_slide_info": [],
        "warnings": [],
    }

    if not path.exists():
        result["warnings"].append(f"File does not exist: {path}")
        return result

    result["file_size_bytes"] = path.stat().st_size

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        result["warnings"].append(f"Failed to open PPTX: {exc}")
        return result

    slides = prs.slides
    result["slide_count"] = len(slides)

    if expected_slide_count is not None and len(slides) != expected_slide_count:
        result["warnings"].append(
            f"Expected {expected_slide_count} slides but found {len(slides)}"
        )

    total_shapes = 0
    per_slide: list[dict] = []

    for idx, slide in enumerate(slides, start=1):
        shape_count = len(slide.shapes)
        total_shapes += shape_count

        slide_info: dict = {
            "slide_number": idx,
            "shape_count": shape_count,
            "shape_names": [],
            "has_text": False,
            "has_image": False,
            "has_chart": False,
            "text_snippets": [],
        }

        for shape in slide.shapes:
            slide_info["shape_names"].append(shape.name or "(unnamed)")

            if shape.has_text_frame:
                texts = [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
                if texts:
                    slide_info["has_text"] = True
                    # Keep first 3 snippets, truncated to 80 chars each
                    for t in texts[:3]:
                        slide_info["text_snippets"].append(t[:80])

            if hasattr(shape, "image"):
                try:
                    _ = shape.image  # accessing triggers AttributeError if not an image
                    slide_info["has_image"] = True
                except Exception:
                    pass

            if shape.has_chart:
                slide_info["has_chart"] = True

        if shape_count == 0:
            result["warnings"].append(f"Slide {idx} has no shapes (empty slide)")

        per_slide.append(slide_info)

    result["total_shapes"] = total_shapes
    result["per_slide_info"] = per_slide

    return result
