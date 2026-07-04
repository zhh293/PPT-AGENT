"""Tests for the template-driven PPT generation pipeline.

Covers:
- OCR module: region detection, line merging, zone classification
- Template ingestion: zone extraction from shapes, index update
- Template index: loading, resolving, slide image paths
- Batch config: full-page background generation with reference images
- PPT writer: background image mode + text overlay
- SlideContent: template_image field
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from ppt_agent.models.slide_contents import SlideContent, SlideZoneContent
from ppt_agent.models.template_meta import (
    build_template_meta_from_ingest,
    default_template_meta,
)
from ppt_agent.retrieval.template_index import load_template_index
from ppt_agent.skills.adapters.gptimage2 import slide_contents_to_batch_config
from ppt_agent.tools.ocr import OCRRegion, _merge_into_lines, classify_zones


# ── OCR tests ──

class TestOCRRegion:
    def test_to_dict(self) -> None:
        r = OCRRegion(text="Hello", x=0.1, y=0.2, w=0.3, h=0.04, confidence=95.0)
        d = r.to_dict()
        assert d["text"] == "Hello"
        assert d["x"] == 0.1
        assert d["confidence"] == 95.0

    def test_merge_into_lines(self) -> None:
        words = [
            OCRRegion("Hello", 0.1, 0.1, 0.1, 0.03, 90.0),
            OCRRegion("World", 0.25, 0.105, 0.1, 0.03, 85.0),
            OCRRegion("Next", 0.1, 0.5, 0.08, 0.03, 88.0),
            OCRRegion("Line", 0.2, 0.505, 0.08, 0.03, 92.0),
        ]
        lines = _merge_into_lines(words, y_tolerance=0.02)
        assert len(lines) == 2
        assert "Hello" in lines[0].text
        assert "World" in lines[0].text
        assert "Next" in lines[1].text

    def test_merge_empty(self) -> None:
        assert _merge_into_lines([]) == []

    def test_classify_zones_title(self) -> None:
        regions = [
            OCRRegion("Big Title", 0.1, 0.05, 0.8, 0.08, 95.0),
            OCRRegion("Body text here", 0.1, 0.4, 0.6, 0.03, 90.0),
            OCRRegion("Page 1", 0.8, 0.92, 0.1, 0.02, 80.0),
        ]
        zones = classify_zones(regions)
        types = {z["type"] for z in zones}
        assert "title" in types
        assert "body" in types
        assert "footer" in types

    def test_classify_zones_empty(self) -> None:
        assert classify_zones([]) == []


# ── Template meta tests ──

class TestTemplateMeta:
    def test_default_has_schema_required_keys(self) -> None:
        meta = default_template_meta(3)
        assert meta["template_id"] == "fallback.default"
        assert meta["slide_count"] == 3
        assert isinstance(meta["color_scheme"], dict)
        assert "primary" in meta["color_scheme"]
        assert len(meta["slides"]) == 3
        for slide in meta["slides"]:
            assert "index" in slide
            assert "layout" in slide
            assert "zones" in slide

    def test_build_from_ingest(self) -> None:
        ingest_meta = {
            "template_id": "blue-medical.report",
            "domain_tags": ["medical"],
            "tone_tags": ["professional"],
            "color_scheme": "blue",
            "theme_colors": {"primary": "#0066CC"},
            "style": "Blue medical template",
            "slide_count": 2,
            "slides": [
                {
                    "index": 0,
                    "zones": [
                        {"zone_id": "t0", "type": "title", "position": [0.1, 0.05, 0.8, 0.1], "text": "Title"},
                        {"zone_id": "b0", "type": "body", "position": [0.1, 0.2, 0.8, 0.6], "text": "Body"},
                        {"zone_id": "i0", "type": "image", "position": [0.7, 0.3, 0.25, 0.4], "text": ""},
                    ],
                    "image_path": "slides/slide_00.png",
                },
                {
                    "index": 1,
                    "zones": [
                        {"zone_id": "t1", "type": "title", "position": [0.1, 0.05, 0.8, 0.1], "text": "Title"},
                    ],
                    "image_path": "slides/slide_01.png",
                },
            ],
        }
        meta = build_template_meta_from_ingest(ingest_meta)
        assert meta["template_id"] == "blue-medical.report"
        assert meta["slide_count"] == 2
        assert len(meta["slides"][0]["text_zones"]) == 2  # title + body
        assert len(meta["slides"][0]["image_zones"]) == 1


# ── Template index tests ──

class TestTemplateIndex:
    def test_load_index_default(self) -> None:
        templates = load_template_index()
        assert isinstance(templates, list)
        # Should have at least the fallback entry
        assert len(templates) >= 1

    def test_load_index_missing_file(self, tmp_path: Path) -> None:
        result = load_template_index(tmp_path / "nonexistent.json")
        assert result == []


# ── Batch config tests ──

class TestBatchConfig:
    def test_full_page_mode(self) -> None:
        slide_contents = {
            "slides": [
                {
                    "slide_index": 0,
                    "template_image": "/path/to/template/slide_00.png",
                    "zones": [
                        {"type": "title", "content": "Introduction"},
                        {"type": "bullets", "content": ["Point A", "Point B"]},
                    ],
                },
                {
                    "slide_index": 1,
                    "zones": [
                        {"type": "title", "content": "Details"},
                    ],
                },
            ]
        }
        config = slide_contents_to_batch_config(slide_contents)
        slides = config["slides"]
        assert len(slides) == 2

        # Slide 0 has reference image (img2img mode)
        assert slides[0]["mode"] == "full_page"
        assert slides[0]["reference_image"] == "/path/to/template/slide_00.png"
        assert "Introduction" in slides[0]["prompt"]
        assert slides[0]["aspect_ratio"] == "16:9"

        # Slide 1 has no reference (text2img mode)
        assert slides[1]["reference_image"] is None
        assert "Details" in slides[1]["prompt"]

    def test_prompt_includes_bullets(self) -> None:
        slide_contents = {
            "slides": [{
                "slide_index": 0,
                "zones": [
                    {"type": "title", "content": "Overview"},
                    {"type": "bullets", "content": ["Feature 1", "Feature 2", "Feature 3"]},
                ],
            }]
        }
        config = slide_contents_to_batch_config(slide_contents)
        prompt = config["slides"][0]["prompt"]
        assert "Overview" in prompt
        assert "Feature 1" in prompt


# ── SlideContent model tests ──

class TestSlideContentModel:
    def test_template_image_field(self) -> None:
        zones = [
            SlideZoneContent("title", "title", [0.1, 0.1, 0.8, 0.1], True, "Hello"),
        ]
        sc = SlideContent(
            slide_index=0,
            layout="content.text",
            zones=zones,
            template_image="/path/to/slide_00.png",
        )
        d = sc.to_dict()
        assert d["template_image"] == "/path/to/slide_00.png"
        assert d["slide_index"] == 0

    def test_to_dict_without_template_image(self) -> None:
        zones = [
            SlideZoneContent("title", "title", [0.1, 0.1, 0.8, 0.1], True, "Hello"),
        ]
        sc = SlideContent(slide_index=0, layout="cover.hero", zones=zones)
        d = sc.to_dict()
        assert d["template_image"] is None


# ── PPT writer tests ──

class TestPPTWriter:
    def test_write_pptx_with_no_background(self, tmp_path: Path) -> None:
        """Without background images, falls back to solid-color mode."""
        from ppt_agent.assembly.ppt_writer import write_pptx

        slide_contents = {
            "slides": [
                {
                    "slide_index": 0,
                    "zones": [
                        {"type": "title", "content": "Test Title", "position": [0.08, 0.08, 0.84, 0.16]},
                        {"type": "bullets", "content": ["A", "B", "C"], "position": [0.1, 0.3, 0.8, 0.5]},
                    ],
                },
            ],
        }
        output = tmp_path / "test.pptx"
        write_pptx(slide_contents, output)
        assert output.exists()
        assert output.stat().st_size > 0

        # Verify content via python-pptx
        from pptx import Presentation
        prs = Presentation(str(output))
        assert len(prs.slides) == 1

    def test_write_pptx_with_background_image(self, tmp_path: Path) -> None:
        """With a background image, uses background + text overlay mode."""
        from PIL import Image
        from ppt_agent.assembly.ppt_writer import write_pptx

        # Create a fake background image
        bg_dir = tmp_path / "background_images"
        bg_dir.mkdir()
        bg_img = bg_dir / "slide_00.png"
        Image.new("RGB", (1920, 1080), color=(30, 60, 120)).save(str(bg_img))

        slide_contents = {
            "slides": [
                {
                    "slide_index": 0,
                    "background_image": str(bg_img),
                    "zones": [
                        {"type": "title", "content": "Title on BG", "position": [0.08, 0.08, 0.84, 0.16]},
                        {"type": "bullets", "content": ["X", "Y"], "position": [0.1, 0.3, 0.5, 0.5]},
                    ],
                },
            ],
        }
        output = tmp_path / "bg_test.pptx"
        write_pptx(slide_contents, output, workspace_root=tmp_path)
        assert output.exists()

        from pptx import Presentation
        prs = Presentation(str(output))
        slide = prs.slides[0]
        # Should have at least 3 shapes: background image + 2 text boxes
        assert len(slide.shapes) >= 3

    def test_write_pptx_auto_detects_background(self, tmp_path: Path) -> None:
        """Test that ppt_writer auto-detects background images from workspace."""
        from PIL import Image
        from ppt_agent.assembly.ppt_writer import write_pptx

        bg_dir = tmp_path / "background_images"
        bg_dir.mkdir()
        Image.new("RGB", (1920, 1080), color=(50, 100, 150)).save(str(bg_dir / "slide_00.png"))

        slide_contents = {
            "slides": [
                {
                    "slide_index": 0,
                    "zones": [
                        {"type": "title", "content": "Auto BG", "position": [0.1, 0.1, 0.8, 0.1]},
                    ],
                },
            ],
        }
        output = tmp_path / "auto_bg.pptx"
        write_pptx(slide_contents, output, workspace_root=tmp_path)
        assert output.exists()

        from pptx import Presentation
        prs = Presentation(str(output))
        # Should have background image + text box
        assert len(prs.slides[0].shapes) >= 2
