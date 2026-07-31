from __future__ import annotations

import json
from copy import deepcopy

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

from ppt_agent.assembly.ppt_writer import (
    _apply_text_to_shape,
    _inject_zone_ids,
    _outline_background_text_zones,
    _replace_shape_image,
    _shorten_background_copy,
    write_pptx,
    write_pptx_from_mapping,
)
from ppt_agent.templates.ingest import (
    _merge_zone_sources,
    refresh_template_meta,
    _unmatched_ocr_zones,
    _zones_from_pptx_shapes,
)
from ppt_agent.workers.content_mapper import (
    _assess_baked_text_risk,
    _archive_resolved_fallback_flags,
    _build_zones_for_slide,
    _build_mapping_diagnostics,
    _ensure_valid_mapping,
    _ensure_generated_content_image_zone,
    _llm_mapping,
    _mapping_llm_repair_target_count,
    _reconcile_zone_repair_flags,
    _repair_zones_in_micro_batches,
    _select_template_slides,
    _synchronize_layered_text_zones,
    _truncate_overflow,
    _validate_mapping_plan,
    _zone_target_range,
    run as run_content_mapper,
)
from ppt_agent.workers.ppt_assembler import (
    _build_mappings_from_slide_contents,
    _fatal_content_mapping_issues,
    _prepare_best_effort_mapping,
    _strict_content_mapping_issues,
)
from ppt_agent.workers.ppt_assembler import run as run_ppt_assembler
from ppt_agent.workers.template_matcher import (
    _choose_required_real_template,
    _has_confident_template_match,
)
from ppt_agent.workers.ppt_verifier import _strict_assembly_invariants
from ppt_agent.coordinator.capabilities import CONTENT_MAPPING
from ppt_agent.models.artifacts import JobWorkspace


def test_background_copy_respects_small_slot_and_arrow_boundaries() -> None:
    zone = {
        "position": [0.57, 0.27, 0.31, 0.065],
        "formatting": {"font_size_pt": 15},
    }

    result = _shorten_background_copy(
        "多模型三阶段共识决策 → 提升规划可靠性与合理性（满意度≥95%）",
        zone,
    )
    assert "→" not in result
    assert len(result) <= 14
    assert not result.endswith("提")


def test_user_image_without_inventory_requests_one_unique_content_image() -> None:
    slide = {"slide_index": 2, "zones": []}
    template_slide = {
        "image_zones": [
            {
                "zone_id": "background",
                "type": "image",
                "position": [0.0, 0.0, 1.0, 1.0],
            },
            {
                "zone_id": "content-frame",
                "type": "image",
                "position": [0.58, 0.28, 0.30, 0.40],
            },
        ]
    }
    design = {
        "visual_strategy": "user_image",
        "block_plan": [{"block_type": "image", "purpose": "AI路线规划界面"}],
    }
    outline = {
        "title": "智能规划",
        "bullets": ["结合用户偏好生成路线"],
    }

    _ensure_generated_content_image_zone(
        slide, template_slide, design, outline, []
    )
    _ensure_generated_content_image_zone(
        slide, template_slide, design, outline, []
    )

    assert len(slide["zones"]) == 1
    zone = slide["zones"][0]
    assert zone["zone_id"] == "content-frame"
    assert zone["action"] == "replace_image"
    assert "AI路线规划界面" in zone["image_prompt"]


def test_generated_background_innovation_uses_semantic_right_rail() -> None:
    slide_data = {
        "slide_index": 3,
        "zones": [
            {
                "zone_id": "old-template-zone",
                "type": "body",
                "position": [0.02, 0.75, 0.85, 0.2],
            }
        ],
    }
    outline_slide = {
        "slide_index": 3,
        "type": "innovation",
        "title": "技术创新矩阵",
        "bullets": [f"创新点{i}" for i in range(5)],
    }

    zones = _outline_background_text_zones(slide_data, outline_slide)

    assert len(zones) == 6
    assert zones[0]["content"] == "技术创新矩阵"
    assert all(zone["position"][0] == 0.57 for zone in zones[1:])
    assert [zone["content"] for zone in zones[1:]] == [
        "创新点0", "创新点1", "创新点2", "创新点3", "创新点4"
    ]


def test_ingest_uses_native_shape_identity() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    shape.text = "Template heading"

    zones = _zones_from_pptx_shapes(slide, 0, prs.slide_width, prs.slide_height)

    assert len(zones) == 1
    assert zones[0]["zone_id"] == f"slide-1/shape-{shape.shape_id}"
    assert zones[0]["native_shape_id"] == shape.shape_id
    assert zones[0]["original_text"] == "Template heading"
    assert zones[0]["legacy_zone_id"] == "s0_shape0"


def test_ingest_recurses_group_shapes_and_records_direction_constraints() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    nested = group.shapes.add_textbox(Inches(2), Inches(2), Inches(3), Inches(1))
    nested.text = "Grouped template copy"
    vertical = group.shapes.add_textbox(Inches(6), Inches(1), Inches(0.4), Inches(3))
    vertical.text = "指标"
    vertical.rotation = 270

    zones = _zones_from_pptx_shapes(slide, 0, prs.slide_width, prs.slide_height)

    grouped = next(zone for zone in zones if zone["native_shape_id"] == nested.shape_id)
    assert grouped["shape_path"] == (
        f"slide-1/group-{group.shape_id}/shape-{nested.shape_id}"
    )
    assert grouped["shape_depth"] == 1
    assert grouped["parent_group_ids"] == [group.shape_id]
    constrained = next(zone for zone in zones if zone["native_shape_id"] == vertical.shape_id)
    assert constrained["shape_rotation_deg"] == 270.0
    assert constrained["content_eligibility"] == "short_label"
    assert constrained["supports_long_text"] is False


def test_ingest_keeps_template_noise_as_actionable_xml_zone() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    footer = slide.shapes.add_textbox(Inches(1), Inches(6.8), Inches(3), Inches(0.25))
    footer.text = "2025-2028年"

    zones = _zones_from_pptx_shapes(slide, 0, prs.slide_width, prs.slide_height)

    assert len(zones) == 1
    assert zones[0]["semantic_role"] == "template_noise"
    assert zones[0]["editable"] is True


def test_ingest_excludes_empty_decorative_autoshapes_but_keeps_empty_textboxes() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(1), Inches(1))
    textbox = slide.shapes.add_textbox(Inches(3), Inches(1), Inches(3), Inches(1))

    zones = _zones_from_pptx_shapes(slide, 0, prs.slide_width, prs.slide_height)

    assert [zone["native_shape_id"] for zone in zones] == [textbox.shape_id]
    assert zones[0]["editable"] is True


def test_ocr_is_audit_only_and_never_becomes_writable_zone() -> None:
    xml_zone = {
        "zone_id": "slide-1/shape-2", "position": [0.1, 0.1, 0.3, 0.1],
        "text": "", "editable": True, "source": "pptx_xml",
    }
    overlapping_ocr = {"position": [0.1, 0.1, 0.3, 0.1], "text": "Rendered"}
    baked_ocr = {"position": [0.7, 0.7, 0.2, 0.1], "text": "In picture"}

    merged = _merge_zone_sources([overlapping_ocr, baked_ocr], [xml_zone])
    baked = _unmatched_ocr_zones([overlapping_ocr, baked_ocr], [xml_zone])

    assert merged == [xml_zone]
    assert merged[0]["ocr_text"] == "Rendered"
    assert baked == [{
        "position": [0.7, 0.7, 0.2, 0.1],
        "text": "In picture",
        "source": "ocr_audit",
        "editable": False,
    }]


def test_refresh_template_meta_is_non_destructive_and_versioned(tmp_path) -> None:
    template_dir = tmp_path / "template"
    template_dir.mkdir()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Nested"
    prs.save(template_dir / "template.pptx")
    (template_dir / "meta.json").write_text(
        json.dumps({"template_id": "test.real", "style": "test", "slides": []}),
        encoding="utf-8",
    )

    meta = refresh_template_meta(template_dir)

    assert (template_dir / "template.pptx").exists()
    assert meta["meta_schema_version"] == "2.0"
    assert meta["parser_strategy"] == "pptx_xml_recursive"
    assert len(meta["template_sha256"]) == 64
    assert meta["slides"][0]["zones"][0]["shape_depth"] == 1


def test_fallback_actions_every_text_zone_without_three_zone_limit() -> None:
    text_zones = [
        {"zone_id": "title", "type": "title", "position": [0.1, 0.1, 0.8, 0.1]},
        {"zone_id": "card-1", "type": "body", "position": [0.1, 0.3, 0.2, 0.3]},
        {"zone_id": "card-2", "type": "body", "position": [0.4, 0.3, 0.2, 0.3]},
        {"zone_id": "card-3", "type": "body", "position": [0.7, 0.3, 0.2, 0.3]},
        {"zone_id": "footer", "type": "footer", "position": [0.1, 0.9, 0.8, 0.05]},
    ]
    zones = _build_zones_for_slide(
        {"title": "Capabilities", "bullets": ["A", "B", "C", "D"]},
        {"text_zones": text_zones, "image_zones": []},
        None,
        0,
    )

    assert len(zones) == 5
    assert {zone.zone_id for zone in zones} == {zone["zone_id"] for zone in text_zones}
    assert sum(zone.action == "replace_text" for zone in zones) == 5
    assert next(zone for zone in zones if zone.zone_id == "footer").content


def test_fallback_never_hard_truncates_complete_source_phrases() -> None:
    text_zones = [
        {
            "zone_id": "title", "type": "title", "position": [0.1, 0.1, 0.8, 0.1],
            "original_text": "项目核心价值", "editable": True,
            "content_eligibility": "title", "max_chars_hint": 10,
        },
        {
            "zone_id": "label", "type": "body", "position": [0.1, 0.3, 0.12, 0.05],
            "original_text": "关键能力", "editable": True,
            "content_eligibility": "short_label", "max_chars_hint": 6,
        },
        {
            "zone_id": "body", "type": "body", "position": [0.25, 0.3, 0.6, 0.2],
            "original_text": "这里是一段用于说明核心方案价值的模板正文", "editable": True,
            "content_eligibility": "body", "supports_long_text": True,
            "max_chars_hint": 30,
        },
    ]

    zones = _build_zones_for_slide(
        {
            "title": "智能借阅提升管理效率",
            "bullets": ["自动登记借还记录", "实时统计库存与热门图书"],
            "source_refs": ["source-1"],
        },
        {"text_zones": text_zones, "image_zones": []},
        None,
        0,
    )

    assert len(zones) == 3
    assert all(zone.action in {"replace_text", "preserve"} for zone in zones)
    assert all(zone.action == "preserve" or str(zone.content).strip() for zone in zones)
    for zone, template_zone in zip(zones, text_zones):
        if zone.action != "replace_text":
            continue
        mapped_length = len(str(zone.content))
        if mapped_length > template_zone["max_chars_hint"]:
            assert zone.fit_status == "overflow"
        assert str(zone.content) in {
            "智能借阅提升管理效率",
            "自动登记借还记录",
            "实时统计库存与热门图书",
            "自动登记借还记录；实时统计库存与热门图书",
        }


def test_capacity_audit_marks_overflow_without_slicing_content() -> None:
    original = "WebSocket实时推送结果"
    zones = [{
        "zone_id": "label", "type": "body", "action": "replace_text",
        "content": original, "fit_status": "unknown", "formatting": {"font_size_pt": 18},
    }]
    template = {"all_zones": [{
        "zone_id": "label", "type": "body", "position": [0.1, 0.1, 0.12, 0.05],
        "original_text": "实时推送", "max_chars_hint": 5,
    }]}

    _truncate_overflow(zones, template)

    assert zones[0]["content"] == original
    assert zones[0]["fit_status"] == "overflow"


def test_capacity_audit_does_not_call_short_copy_overflow() -> None:
    zones = [{
        "zone_id": "body", "type": "body", "action": "replace_text",
        "content": "缓存设计差", "fit_status": "unknown",
        "formatting": {"font_size_pt": 11.5},
    }]
    template_zone = {
        "zone_id": "body", "type": "body",
        "position": [0.1, 0.4, 0.2232, 0.1013],
        "original_text": "走访行业头部企业，分析领先技术方案" * 3,
        "max_chars_hint": 24,
    }

    lower, upper = _zone_target_range(template_zone)
    _truncate_overflow(zones, {"all_zones": [template_zone]})

    assert (lower, upper) == (13, 24)
    assert zones[0]["fit_status"] == "fits"
    assert _validate_mapping_plan(
        {"zones": zones},
        {"text_zones": [template_zone], "all_zones": [template_zone]},
    ) == ["length_mismatch:body"]


def test_capacity_cap_does_not_require_exact_short_label_length() -> None:
    lower, upper = _zone_target_range({
        "type": "body",
        "content_eligibility": "short_label",
        "original_text": "2222",
        "max_chars_hint": 2,
        "position": [0.7, 0.5, 0.0635, 0.0583],
        "formatting": {"font_size_pt": 20},
    })

    assert (lower, upper) == (1, 2)


def test_mapping_validator_rejects_clear_or_empty_text_and_length_mismatch() -> None:
    template_zone = {
        "zone_id": "required", "type": "body", "editable": True,
        "original_text": "长度接近模板文字", "max_chars_hint": 12,
        "content_eligibility": "body",
    }

    clear_issues = _validate_mapping_plan(
        {"zones": [{"zone_id": "required", "action": "clear_text", "content": None}]},
        {"text_zones": [template_zone], "all_zones": [template_zone]},
    )
    long_issues = _validate_mapping_plan(
        {"zones": [{
            "zone_id": "required", "action": "replace_text",
            "content": "这是一段明显远远超过模板原文字数限制的长文本内容",
            "fit_status": "fits",
        }]},
        {"text_zones": [template_zone], "all_zones": [template_zone]},
    )

    assert "clear_text_forbidden:required" in clear_issues
    assert "length_mismatch:required" in long_issues


def test_overlapping_duplicate_template_layers_receive_identical_copy() -> None:
    template_zones = [
        {"zone_id": "shadow", "original_text": "核心能力", "position": [0.1, 0.2, 0.3, 0.1]},
        {"zone_id": "foreground", "original_text": "核心能力", "position": [0.102, 0.201, 0.3, 0.1]},
    ]
    zones = [
        {"zone_id": "shadow", "action": "replace_text", "content": "智能借阅"},
        {"zone_id": "foreground", "action": "replace_text", "content": "库存管理"},
    ]

    _synchronize_layered_text_zones(zones, {"text_zones": template_zones})

    assert zones[0]["content"] == zones[1]["content"]


def test_wide_overlapping_duplicate_layers_receive_identical_copy() -> None:
    template_zones = [
        {
            "zone_id": "outline",
            "original_text": "用户价值与行业影响",
            "position": [0.04, 0.18, 0.91, 0.07],
        },
        {
            "zone_id": "foreground",
            "original_text": "用户价值与行业影响",
            "position": [0.127, 0.179, 0.733, 0.069],
        },
    ]
    zones = [
        {"zone_id": "outline", "action": "replace_text", "content": "用户价值"},
        {"zone_id": "foreground", "action": "replace_text", "content": "行业影响"},
    ]

    _synchronize_layered_text_zones(zones, {"text_zones": template_zones})

    assert zones[0]["content"] == zones[1]["content"]


def test_mapping_diagnostics_include_legacy_text_zones_without_editable_flag() -> None:
    template_zone = {
        "zone_id": "s0_shape1", "type": "title",
        "text": "模板标题", "position": [0.1, 0.1, 0.5, 0.1],
    }
    payload = {"slides": [{
        "slide_index": 0, "template_slide_index": 0,
        "zones": [{
            "zone_id": "s0_shape1", "type": "title", "editable": True,
            "action": "replace_text", "content": "项目标题", "fit_status": "fits",
        }],
    }]}

    diagnostics = _build_mapping_diagnostics(payload, {
        "slides": [{
            "index": 0, "text_zones": [template_zone],
            "all_zones": [template_zone],
        }]
    })

    assert len(diagnostics["slides"][0]["zones"]) == 1


def test_llm_mapping_fallback_batch_is_marked_for_review() -> None:
    class FallbackClient:
        def generate_json(self, **kwargs):
            return kwargs["fallback"]

        def was_fallback(self, phase: str) -> bool:
            return phase == "content_mapping_0"

    template_zone = {
        "zone_id": "title", "type": "title", "editable": True,
        "position": [0.1, 0.1, 0.7, 0.1], "original_text": "模板标题",
        "content_eligibility": "title", "max_chars_hint": 8,
    }
    result = _llm_mapping(
        FallbackClient(),
        {"meta": {}, "slides": [{
            "slide_index": 0, "title": "项目标题", "bullets": ["项目内容"],
            "source_refs": ["source-1"],
        }]},
        {"template_id": "real.template"},
        {"theme_profile": {}, "slides": [{"slide_index": 0}]},
        {"image_inventory": []},
        {"slides": [{
            "index": 0, "layout": "cover", "text_zones": [template_zone],
            "image_zones": [], "all_zones": [template_zone],
        }]},
    )

    assert result["slides"][0]["review_status"] == "needs_review"
    assert "llm_mapping_fallback:batch_0" in result["slides"][0]["fallback_flags"]


def test_incomplete_llm_batch_recovers_missing_slide_without_reindexing() -> None:
    class RecoveringClient:
        def __init__(self) -> None:
            self.phases: list[str] = []

        def generate_json(self, **kwargs):
            phase = kwargs["phase"]
            self.phases.append(phase)
            if phase == "content_mapping_0":
                return {
                    "template_id": "real.template",
                    "slides": [{
                        "slide_index": 0, "template_slide_index": 0,
                        "zones": [{
                            "zone_id": "s0-title", "type": "title",
                            "content": "第一页", "action": "replace_text",
                            "placement_reason": "first slide title",
                            "source_block_ids": ["slide-0"], "transformation": "rewrite",
                        }],
                    }],
                }
            if phase == "content_mapping_retry_1_1":
                return {
                    "template_id": "real.template",
                    "slides": [{
                        "slide_index": 1, "template_slide_index": 1,
                        "zones": [{
                            "zone_id": "s1-title", "type": "title",
                            "content": "第二页", "action": "replace_text",
                            "placement_reason": "recovered second slide title",
                            "source_block_ids": ["slide-1"], "transformation": "rewrite",
                        }],
                    }],
                }
            return kwargs["fallback"]

        def was_fallback(self, phase: str) -> bool:
            return False

    def tpl(index: int) -> dict:
        zone = {
            "zone_id": f"s{index}-title", "type": "title", "editable": True,
            "position": [0.1, 0.1, 0.7, 0.1], "original_text": "模板标题",
            "content_eligibility": "title", "max_chars_hint": 8,
        }
        return {
            "index": index, "layout": "content.title", "text_zones": [zone],
            "image_zones": [], "all_zones": [zone],
        }

    client = RecoveringClient()
    result = _llm_mapping(
        client,
        {"meta": {}, "slides": [
            {"slide_index": 0, "title": "第一页", "bullets": [], "source_refs": ["slide-0"]},
            {"slide_index": 1, "title": "第二页", "bullets": [], "source_refs": ["slide-1"]},
        ]},
        {"template_id": "real.template"},
        {"theme_profile": {}, "slides": [{"slide_index": 0}, {"slide_index": 1}]},
        {"image_inventory": []},
        {"slides": [tpl(0), tpl(1)]},
    )

    assert [slide["slide_index"] for slide in result["slides"]] == [0, 1]
    assert "content_mapping_retry_1_1" in client.phases
    assert "llm_batch_recovered:batch_0" in result["slides"][1]["fallback_flags"]
    assert not any(
        flag.startswith("deterministic_last_resort")
        for slide in result["slides"] for flag in slide.get("fallback_flags", [])
    )


def test_out_of_range_zone_is_semantically_rewritten_by_llm() -> None:
    class FittingClient:
        def __init__(self) -> None:
            self.phases: list[str] = []

        def generate_json(self, **kwargs):
            phase = kwargs["phase"]
            self.phases.append(phase)
            if phase == "content_mapping_0":
                return {
                    "template_id": "real.template",
                    "slides": [{
                        "slide_index": 0, "template_slide_index": 0,
                        "zones": [{
                            "zone_id": "title", "type": "title",
                            "content": "这是一个明显超过文本框容量的完整标题",
                            "action": "replace_text", "placement_reason": "semantic title",
                            "source_block_ids": ["slide0_title"],
                            "transformation": "rewrite", "fit_status": "unknown",
                        }],
                    }],
                }
            if phase == "content_mapping_micro_0_1":
                return {"repairs": [{
                    "zone_id": "title", "content": "智能规划",
                    "placement_reason": "semantic title",
                    "source_block_ids": ["slide0_title"],
                    "transformation": "rewrite",
                }]}
            return kwargs["fallback"]

        def was_fallback(self, phase: str) -> bool:
            return False

    zone = {
        "zone_id": "title", "type": "title", "editable": True,
        "position": [0.1, 0.1, 0.8, 0.1], "original_text": "核心能力",
        "content_eligibility": "title", "max_chars_hint": 6,
    }
    client = FittingClient()
    result = _llm_mapping(
        client,
        {"meta": {}, "slides": [{
            "slide_index": 0, "title": "智能规划", "bullets": [],
            "source_refs": ["slide-0"],
        }]},
        {"template_id": "real.template"},
        {"theme_profile": {}, "slides": [{"slide_index": 0}]},
        {"image_inventory": []},
        {"slides": [{
            "index": 0, "layout": "cover", "text_zones": [zone],
            "image_zones": [], "all_zones": [zone],
        }]},
    )

    mapped = result["slides"][0]["zones"][0]
    assert mapped["content"] == "智能规划"
    assert mapped["source"] == "llm"
    assert mapped["fit_status"] == "fits"
    assert result["slides"][0]["llm_repair_count"] == 1
    assert "content_mapping_micro_0_1" in client.phases


def test_incomplete_slide_fit_pass_uses_micro_zone_repairs() -> None:
    class MicroRepairClient:
        def __init__(self) -> None:
            self.phases: list[str] = []

        def generate_json(self, **kwargs):
            phase = kwargs["phase"]
            self.phases.append(phase)
            if phase == "content_mapping_0":
                return {"template_id": "real.template", "slides": []}
            if phase.startswith("content_mapping_fit_"):
                return {"template_id": "real.template", "slides": []}
            if phase == "content_mapping_micro_0_1":
                return {"repairs": [
                    {
                        "zone_id": "label-a", "content": "智能",
                        "placement_reason": "concise capability label",
                        "source_block_ids": ["slide0_title"],
                        "transformation": "rewrite",
                    },
                    {
                        "zone_id": "label-b", "content": "规划",
                        "placement_reason": "concise capability label",
                        "source_block_ids": ["slide0_title"],
                        "transformation": "rewrite",
                    },
                ]}
            return kwargs["fallback"]

        def was_fallback(self, phase: str) -> bool:
            return False

    zones = [
        {
            "zone_id": zone_id, "type": "title", "editable": True,
            "position": [0.1 + index * 0.3, 0.1, 0.25, 0.1],
            "original_text": "模板甲", "content_eligibility": "title",
            "max_chars_hint": 4,
        }
        for index, zone_id in enumerate(("label-a", "label-b"))
    ]
    client = MicroRepairClient()
    result = _llm_mapping(
        client,
        {"meta": {}, "slides": [{
            "slide_index": 0, "title": "智能规划", "bullets": [],
            "source_refs": ["slide0_title"],
        }]},
        {"template_id": "real.template"},
        {"theme_profile": {}, "slides": [{"slide_index": 0}]},
        {"image_inventory": []},
        {"slides": [{
            "index": 0, "layout": "content.labels",
            "text_zones": zones, "image_zones": [], "all_zones": zones,
        }]},
    )

    mapped = {zone["zone_id"]: zone for zone in result["slides"][0]["zones"]}
    assert mapped["label-a"]["content"] == "智能"
    assert mapped["label-b"]["content"] == "规划"
    assert all(mapped[zone_id]["source"] == "llm" for zone_id in mapped)
    assert all(mapped[zone_id]["fit_status"] == "fits" for zone_id in mapped)
    assert result["slides"][0]["llm_micro_repair_count"] == 2
    assert "content_mapping_micro_0_1" in client.phases
    assert "content_mapping_retry_0_1" in client.phases
    assert not any(
        flag.startswith(("missing_llm_slide:", "llm_mapping_fallback:"))
        for flag in result["slides"][0]["fallback_flags"]
    )


def test_invalid_micro_batch_copy_retries_only_rejected_zone() -> None:
    class RetryingMicroClient:
        def __init__(self) -> None:
            self.phases: list[str] = []

        def generate_json(self, **kwargs):
            phase = kwargs["phase"]
            self.phases.append(phase)
            if phase == "content_mapping_0":
                return {
                    "template_id": "real.template",
                    "slides": [{
                        "slide_index": 0, "template_slide_index": 0,
                        "zones": [{
                            "zone_id": "title", "type": "title",
                            "content": "这是明显过长的标题", "action": "replace_text",
                            "placement_reason": "initial mapping",
                            "source_block_ids": ["slide0_title"],
                            "transformation": "rewrite", "fit_status": "unknown",
                        }],
                    }],
                }
            if phase.startswith("content_mapping_fit_"):
                return {"template_id": "real.template", "slides": []}
            content = (
                "仍然明显过长"
                if phase == "content_mapping_micro_0_1"
                else "智能"
            )
            return {"repairs": [{
                "zone_id": "title", "content": content,
                "placement_reason": "semantic rewrite",
                "source_block_ids": ["slide0_title"],
                "transformation": "rewrite",
            }]}

        def was_fallback(self, phase: str) -> bool:
            return False

    zone = {
        "zone_id": "title", "type": "title", "editable": True,
        "position": [0.1, 0.1, 0.4, 0.1], "original_text": "模板甲",
        "content_eligibility": "title", "max_chars_hint": 4,
    }
    client = RetryingMicroClient()
    result = _llm_mapping(
        client,
        {"meta": {}, "slides": [{
            "slide_index": 0, "title": "智能规划", "bullets": [],
            "source_refs": ["slide0_title"],
        }]},
        {"template_id": "real.template"},
        {"theme_profile": {}, "slides": [{"slide_index": 0}]},
        {"image_inventory": []},
        {"slides": [{
            "index": 0, "layout": "content.title",
            "text_zones": [zone], "image_zones": [], "all_zones": [zone],
        }]},
    )

    mapped = result["slides"][0]["zones"][0]
    assert mapped["content"] == "智能"
    assert mapped["fit_status"] == "fits"
    assert "content_mapping_micro_0_1_retry_1" in client.phases


def test_draft_cached_mapping_is_targeted_repaired_instead_of_returned(tmp_path) -> None:
    class CachedRepairClient:
        def __init__(self) -> None:
            self.phases: list[str] = []

        def generate_json(self, **kwargs):
            phase = kwargs["phase"]
            self.phases.append(phase)
            return {"repairs": [{
                "zone_id": "title", "content": "智能",
                "placement_reason": "cached zone semantic rewrite",
                "source_block_ids": ["slide0_title"],
                "transformation": "rewrite",
            }]}

        def was_fallback(self, phase: str) -> bool:
            return False

    zone = {
        "zone_id": "title", "type": "title", "editable": True,
        "position": [0.1, 0.1, 0.4, 0.1], "original_text": "模板甲",
        "content_eligibility": "title", "max_chars_hint": 4,
    }
    artifacts = {
        "outline.json": {"meta": {}, "slides": [{
            "slide_index": 0, "title": "智能规划", "bullets": [],
            "source_refs": ["slide0_title"],
        }]},
        "selected_template.json": {"template_id": "real.template"},
        "slide_design_plan.json": {
            "theme_profile": {}, "slides": [{"slide_index": 0}],
        },
        "source_summary.json": {"image_inventory": []},
        "template_zones.json": {"slides": [{
            "index": 0, "layout": "content.title", "text_zones": [zone],
            "image_zones": [], "all_zones": [zone],
        }]},
        "slide_contents.json": {
            "template_id": "real.template", "review_status": "draft",
            "mapping_mode": "llm_first",
            "slides": [{
                "slide_index": 0, "template_slide_index": 0,
                "review_status": "draft", "fallback_flags": [],
                "zones": [{
                    **zone, "action": "replace_text", "content": "规划",
                    "source": "generated", "fit_status": "fits",
                    "placement_reason": "old deterministic fallback",
                    "source_block_ids": ["slide0_title"],
                    "transformation": "rewrite",
                }],
            }],
        },
    }
    for filename, payload in artifacts.items():
        (tmp_path / filename).write_text(
            json.dumps(payload, ensure_ascii=False), encoding="utf-8"
        )

    client = CachedRepairClient()
    workspace = JobWorkspace(tmp_path)
    output = run_content_mapper(workspace, force=False, llm_client=client)
    repaired = json.loads(output.read_text(encoding="utf-8"))

    assert _mapping_llm_repair_target_count(
        artifacts["slide_contents.json"], artifacts["template_zones.json"]
    ) == 1
    assert client.phases == ["content_mapping_micro_0_1"]
    assert repaired["slides"][0]["zones"][0]["source"] == "llm"
    assert repaired["targeted_repair"] == {
        "mode": "cached_zone_micro_repair",
        "initial_target_zones": 1,
        "remaining_target_zones": 0,
    }
    assert repaired["review_status"] == "draft"


def test_approved_cached_mapping_is_repaired_without_losing_approval(tmp_path) -> None:
    class RepairClient:
        def __init__(self) -> None:
            self.phases: list[str] = []

        def generate_json(self, **kwargs):
            self.phases.append(kwargs["phase"])
            return {"repairs": [{
                "zone_id": "title",
                "content": "Approved title",
                "placement_reason": "approved copy repaired before assembly",
                "source_block_ids": ["slide0_title"],
                "transformation": "rewrite",
            }]}

        def was_fallback(self, phase: str) -> bool:
            return False

    zone = {
        "zone_id": "title",
        "type": "title",
        "editable": True,
        "position": [0.1, 0.1, 0.4, 0.1],
        "original_text": "Template title",
        "content_eligibility": "title",
        "max_chars_hint": 20,
    }
    artifacts = {
        "outline.json": {
            "meta": {},
            "slides": [{
                "slide_index": 0,
                "title": "Approved title",
                "bullets": [],
                "source_refs": [],
            }],
        },
        "selected_template.json": {"template_id": "real.template"},
        "slide_design_plan.json": {
            "theme_profile": {},
            "slides": [{"slide_index": 0}],
        },
        "source_summary.json": {"image_inventory": []},
        "template_zones.json": {
            "slides": [{
                "index": 0,
                "layout": "content.title",
                "text_zones": [zone],
                "image_zones": [],
                "all_zones": [zone],
            }],
        },
        "slide_contents.json": {
            "template_id": "real.template",
            "review_status": "approved",
            "mapping_mode": "llm_first",
            "slides": [{
                "slide_index": 0,
                "template_slide_index": 0,
                "review_status": "approved",
                "fallback_flags": [],
                "zones": [{
                    **zone,
                    "action": "replace_text",
                    "content": "Approved title",
                    "source": "generated",
                    "fit_status": "fits",
                    "placement_reason": "approved copy",
                    "source_block_ids": [],
                    "transformation": "rewrite",
                }],
            }],
        },
    }
    for filename, payload in artifacts.items():
        (tmp_path / filename).write_text(
            json.dumps(payload, ensure_ascii=False),
            encoding="utf-8",
        )

    client = RepairClient()
    output = run_content_mapper(
        JobWorkspace(tmp_path),
        force=False,
        llm_client=client,
    )

    repaired = json.loads(output.read_text(encoding="utf-8"))
    assert client.phases == ["content_mapping_micro_0_1"]
    assert repaired["review_status"] == "approved"
    assert repaired["slides"][0]["zones"][0]["source"] == "llm"
    assert repaired["targeted_repair"]["remaining_target_zones"] == 0


def test_micro_repair_ignores_non_object_items_instead_of_crashing() -> None:
    class InvalidRepairClient:
        def generate_json(self, **kwargs):
            return {"repairs": ["invalid repair item"]}

        def was_fallback(self, phase: str) -> bool:
            return False

    slide = {
        "slide_index": 0,
        "template_slide_index": 0,
        "zones": [{
            "zone_id": "title",
            "type": "title",
            "content": "Old title",
            "action": "replace_text",
            "source": "deterministic",
            "source_block_ids": ["slide-0-title"],
            "fit_status": "fits",
        }],
        "fallback_flags": [],
    }
    outline_slide = {
        "slide_index": 0,
        "title": "Intelligent travel",
        "bullets": ["Automatic itinerary planning"],
    }
    template_slide = {
        "index": 0,
        "text_zones": [{
            "zone_id": "title",
            "type": "title",
            "editable": True,
            "content_eligibility": "title",
            "original_text": "Template title",
            "max_chars_hint": 20,
            "position": [0.1, 0.1, 0.8, 0.1],
            "formatting": {"font_size_pt": 24},
        }],
    }

    repaired = _repair_zones_in_micro_batches(
        InvalidRepairClient(),
        slide,
        outline_slide,
        [],
        {},
        {},
        template_slide,
    )

    assert repaired == 0
    assert slide["zones"][0]["content"] == "Old title"
    assert any(
        flag.startswith("llm_micro_repair_unresolved:")
        for flag in slide["fallback_flags"]
    )


def test_strict_content_mapping_gate_cannot_be_bypassed_by_force() -> None:
    issues = _strict_content_mapping_issues({
        "slides": [{
            "slide_index": 0,
            "zones": [
                {"zone_id": "empty", "type": "body", "editable": True,
                 "action": "clear_text", "content": None},
                {"zone_id": "overflow", "type": "title", "editable": True,
                 "action": "replace_text", "content": "Too long", "fit_status": "overflow"},
            ],
        }],
    })

    assert "slide_0:clear_text_forbidden:empty" in issues
    assert "slide_0:overflow:overflow" in issues


def test_text_fit_issues_are_recoverable_for_best_effort_assembly() -> None:
    payload = {
        "slides": [{
            "slide_index": 0,
            "zones": [
                {
                    "zone_id": "empty",
                    "type": "body",
                    "editable": True,
                    "action": "clear_text",
                    "content": None,
                },
                {
                    "zone_id": "overflow",
                    "type": "title",
                    "editable": True,
                    "action": "replace_text",
                    "content": "Too long",
                    "fit_status": "overflow",
                },
            ],
        }],
    }
    issues = _strict_content_mapping_issues(payload)

    assert _fatal_content_mapping_issues(issues) == []
    prepared, recoveries = _prepare_best_effort_mapping(payload, issues)
    assert prepared["slides"][0]["zones"][0]["action"] == "preserve"
    assert prepared["slides"][0]["zones"][1]["content"] == "Too long"
    assert recoveries[0]["action"] == "clear_text_reverted_to_template"


def test_slide_sequence_mismatch_remains_fatal_for_assembly() -> None:
    issues = _strict_content_mapping_issues(
        {"slides": [{"slide_index": 1, "zones": []}]},
        expected_slide_indices=[0],
    )

    assert _fatal_content_mapping_issues(issues) == issues


def test_reconcile_removes_stale_zone_issue_flags_after_valid_repair() -> None:
    slide = {
        "zones": [{
            "zone_id": "title",
            "type": "title",
            "action": "replace_text",
            "content": "OK",
            "source": "llm",
            "source_block_ids": ["slide0_title"],
            "transformation": "rewrite",
            "fit_status": "fits",
        }],
        "fallback_flags": [
            "length_mismatch:title",
            "untraceable_rewrite:title",
            "overflow:title",
            "llm_micro_repair_unresolved:title",
        ],
    }
    template_slide = {
        "text_zones": [{
            "zone_id": "title",
            "type": "title",
            "editable": True,
            "original_text": "Old title",
            "min_chars_hint": 3,
            "max_chars_hint": 20,
        }]
    }

    _reconcile_zone_repair_flags(slide, template_slide)

    assert slide["fallback_flags"] == []


def test_passed_diagnostics_archive_recovery_history() -> None:
    payload = {
        "slides": [{
            "slide_index": 0,
            "fallback_flags": ["llm_template_slide_override_rejected:1->0"],
        }]
    }

    _archive_resolved_fallback_flags(payload, {"issue_count": 0})

    slide = payload["slides"][0]
    assert slide["fallback_flags"] == []
    assert slide["mapping_history_flags"] == [
        "llm_template_slide_override_rejected:1->0"
    ]


def test_llm_first_mapping_gate_requires_complete_indices_and_high_llm_ratio() -> None:
    issues = _strict_content_mapping_issues(
        {
            "mapping_mode": "llm_first",
            "generation_stats": {
                "passed": False,
                "llm_text_ratio": 0.5,
                "minimum_llm_text_ratio": 0.85,
            },
            "slides": [{"slide_index": 0, "zones": []}],
        },
        expected_slide_indices=[0, 1],
    )

    assert any(issue.startswith("slide_index_mismatch:") for issue in issues)
    assert "llm_text_ratio_below_minimum:0.5<0.85" in issues


def test_multiple_body_zones_do_not_overwrite_each_other() -> None:
    slide_contents = {
        "slides": [{
            "slide_index": 0,
            "zones": [
                {"zone_id": "body-1", "type": "body", "content": "One", "action": "replace_text"},
                {"zone_id": "body-2", "type": "body", "content": "Two", "action": "replace_text"},
                {"zone_id": "body-3", "type": "body", "content": None, "action": "clear_text"},
            ],
        }],
    }

    mappings, _ = _build_mappings_from_slide_contents(slide_contents)

    assert set(mappings[0]) == {"body-1", "body-2", "body-3"}
    assert mappings[0]["body-3"]["action"] == "clear_text"


def test_mapping_plan_validator_rejects_duplicates_overlays_and_missing_actions() -> None:
    template_slide = {
        "text_zones": [
            {"zone_id": "title", "editable": True},
            {"zone_id": "body", "editable": True},
        ],
        "all_zones": [
            {"zone_id": "title", "editable": True},
            {"zone_id": "body", "editable": True},
        ],
    }
    slide = {
        "zones": [
            {"zone_id": "title", "action": "replace_text", "content": "A"},
            {"zone_id": "title", "action": "replace_text", "content": "B"},
            {"zone_id": "_overlay_0_body", "action": "replace_text", "content": "C"},
        ]
    }

    issues = _validate_mapping_plan(slide, template_slide)

    assert "duplicate_zone:title" in issues
    assert "unactioned_zone:body" in issues
    assert "overlay_forbidden:_overlay_0_body" in issues


def test_mapping_plan_rejects_long_text_in_rotated_short_label_zone() -> None:
    constrained = {
        "zone_id": "vertical-label", "editable": True,
        "content_eligibility": "short_label", "max_chars_hint": 6,
    }
    issues = _validate_mapping_plan(
        {"zones": [{
            "zone_id": "vertical-label", "action": "replace_text",
            "content": "This is long prose", "fit_status": "unknown",
        }]},
        {"text_zones": [constrained], "all_zones": [constrained]},
    )

    assert "content_eligibility_violation:vertical-label" in issues


def test_mapping_plan_treats_punctuation_only_copy_as_missing() -> None:
    zone = {
        "zone_id": "metric", "editable": True, "type": "body",
        "original_text": "3200万元", "position": [0.1, 0.1, 0.3, 0.1],
    }

    issues = _validate_mapping_plan(
        {"zones": [{
            "zone_id": "metric", "action": "replace_text",
            "content": "--", "fit_status": "fits",
        }]},
        {"text_zones": [zone], "all_zones": [zone]},
    )

    assert "missing_replacement_text:metric" in issues


def test_native_shape_resolution_and_text_replacement_preserve_formatting() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    paragraph = shape.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = "Original"
    run.font.name = "Arial"
    run.font.size = Pt(23)
    original_geometry = (shape.left, shape.top, shape.width, shape.height)

    mapping = _inject_zone_ids(slide, {
        "all_zones": [{
            "zone_id": "stable-zone",
            "native_shape_id": shape.shape_id,
            "position": [0, 0, 0.01, 0.01],
        }],
    })
    _apply_text_to_shape(mapping["stable-zone"], {"type": "title", "content": "Replacement"})

    replaced_run = shape.text_frame.paragraphs[0].runs[0]
    assert replaced_run.text == "Replacement"
    assert replaced_run.font.name == "Arial"
    assert replaced_run.font.size == Pt(23)
    assert (shape.left, shape.top, shape.width, shape.height) == original_geometry


def test_native_shape_path_resolves_nested_group_child() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    nested = group.shapes.add_textbox(Inches(2), Inches(2), Inches(3), Inches(1))
    nested.text = "Original"
    zone_id = f"slide-1/group-{group.shape_id}/shape-{nested.shape_id}"

    mapping = _inject_zone_ids(slide, {
        "all_zones": [{
            "zone_id": zone_id,
            "shape_path": zone_id,
            "native_shape_id": nested.shape_id,
            "position": [0.2, 0.2, 0.3, 0.1],
        }],
    })

    assert mapping[zone_id].shape_id == nested.shape_id
    assert mapping[zone_id].text == "Original"


def test_baked_text_risk_is_confirmed_when_ocr_is_not_editable() -> None:
    slide = {
        "ocr_text": "Health monitoring platform",
        "text_zones": [{"text": "Editable heading"}],
        "image_zones": [{"position": [0, 0, 1, 1]}],
    }

    risk = _assess_baked_text_risk(slide)

    assert risk["level"] == "confirmed"
    assert risk["unmatched_ocr_ratio"] > 0.5


def test_template_slide_selection_is_unique_ordered_and_prefers_compatible_pages() -> None:
    def template_slide(index: int, text_count: int, *, ocr: str = "") -> dict:
        zones = [
            {
                "zone_id": f"s{index}-{zone_index}",
                "type": "title" if zone_index == 0 else "body",
                "position": [0.1, 0.1 + zone_index * 0.1, 0.8, 0.08],
                "text": "Template copy",
                "editable": True,
            }
            for zone_index in range(text_count)
        ]
        return {
            "index": index,
            "layout": "content.text",
            "text_zones": zones,
            "all_zones": zones,
            "image_zones": [{"position": [0, 0, 1, 1]}] if ocr else [],
            "ocr_text": ocr,
        }

    template = {
        "slides": [
            template_slide(0, 1),
            template_slide(1, 4, ocr="Uneditable medical product wording"),
            template_slide(2, 4),
            template_slide(3, 3),
        ]
    }
    outline = {
        "slides": [
            {"slide_index": 0, "type": "content", "title": "A", "bullets": ["1", "2", "3"]},
            {"slide_index": 1, "type": "content", "title": "B", "bullets": ["1", "2"]},
        ]
    }

    selected = _select_template_slides(outline, template, {})

    assert selected == sorted(set(selected))
    assert len(selected) == 2
    assert 1 not in selected


def test_mapping_rejects_llm_template_page_reordering() -> None:
    template = {
        "slides": [
            {
                "index": index,
                "layout": f"layout-{index}",
                "text_zones": [],
                "all_zones": [],
            }
            for index in range(4)
        ]
    }
    outline = {
        "slides": [
            {"slide_index": 0, "type": "cover", "title": "Cover", "bullets": []},
            {"slide_index": 1, "type": "closing", "title": "Close", "bullets": []},
        ]
    }
    expected = _select_template_slides(outline, template, {})
    mapping = {
        "slides": [
            {"slide_index": 0, "template_slide_index": expected[1], "zones": []},
            {"slide_index": 1, "template_slide_index": expected[0], "zones": []},
        ]
    }

    _ensure_valid_mapping(mapping, outline, {}, {}, {}, template)

    assert [slide["template_slide_index"] for slide in mapping["slides"]] == expected
    assert all(
        any(
            flag.startswith("llm_template_slide_override_rejected:")
            for flag in slide["fallback_flags"]
        )
        for slide in mapping["slides"]
    )


def test_mapping_resets_majority_cross_template_zone_set() -> None:
    def template_slide(index: int) -> dict:
        zones = [
            {
                "zone_id": f"slide-{index}-zone-{zone_index}",
                "type": "title" if zone_index == 0 else "body",
                "editable": True,
                "original_text": f"Template {index} text {zone_index}",
                "position": [0.1, 0.1 + zone_index * 0.2, 0.8, 0.1],
            }
            for zone_index in range(2)
        ]
        return {
            "index": index,
            "layout": f"layout-{index}",
            "text_zones": zones,
            "all_zones": zones,
        }

    template = {"slides": [template_slide(0), template_slide(1)]}
    outline = {
        "slides": [
            {"slide_index": 0, "type": "cover", "title": "Cover", "bullets": []},
            {"slide_index": 1, "type": "closing", "title": "Close", "bullets": []},
        ]
    }
    mapping = {
        "slides": [
            {
                "slide_index": 0,
                "template_slide_index": 0,
                "zones": [
                    {"zone_id": "slide-1-zone-0", "type": "title", "content": "Wrong"},
                    {"zone_id": "slide-1-zone-1", "type": "body", "content": "Wrong"},
                ],
            },
            {
                "slide_index": 1,
                "template_slide_index": 1,
                "zones": [
                    {"zone_id": "slide-0-zone-0", "type": "title", "content": "Wrong"},
                    {"zone_id": "slide-0-zone-1", "type": "body", "content": "Wrong"},
                ],
            },
        ]
    }

    _ensure_valid_mapping(mapping, outline, {}, {}, {}, template)

    for slide in mapping["slides"]:
        prefix = f"slide-{slide['template_slide_index']}-zone-"
        assert all(zone["zone_id"].startswith(prefix) for zone in slide["zones"])
        assert any(
            flag.startswith("cross_template_zone_set_rejected:")
            for flag in slide["fallback_flags"]
        )


def test_mapping_discards_individual_unknown_zone_ids() -> None:
    template_zone = {
        "zone_id": "known",
        "type": "title",
        "editable": True,
        "original_text": "Template title",
        "position": [0.1, 0.1, 0.8, 0.1],
    }
    mapping = {
        "slides": [{
            "slide_index": 0,
            "template_slide_index": 0,
            "zones": [
                {"zone_id": "known", "type": "title", "content": "Mapped title"},
                {"zone_id": "unknown", "type": "body", "content": "Discard me"},
            ],
        }]
    }

    _ensure_valid_mapping(
        mapping,
        {"slides": [{"slide_index": 0, "type": "cover", "title": "Mapped title"}]},
        {},
        {},
        {},
        {"slides": [{
            "index": 0,
            "layout": "cover",
            "text_zones": [template_zone],
            "all_zones": [template_zone],
        }]},
    )

    assert [zone["zone_id"] for zone in mapping["slides"][0]["zones"]] == ["known"]
    assert "invalid_zones_discarded:1" in mapping["slides"][0]["fallback_flags"]


def test_llm_mapping_retries_when_model_swaps_valid_template_pages() -> None:
    class SwappingClient:
        def __init__(self) -> None:
            self.phases = []

        def generate_json(self, **kwargs):
            phase = kwargs["phase"]
            self.phases.append(phase)
            if phase == "content_mapping_0":
                return {
                    "slides": [
                        {"slide_index": 0, "template_slide_index": 1, "zones": []},
                        {"slide_index": 1, "template_slide_index": 0, "zones": []},
                    ]
                }
            requested = kwargs["context"]["expected_slide_indices"][0]
            return {
                "slides": [{
                    "slide_index": requested,
                    "template_slide_index": requested,
                    "zones": [],
                }]
            }

        def was_fallback(self, _phase):
            return False

    outline = {
        "meta": {},
        "slides": [
            {"slide_index": 0, "type": "cover", "title": "Cover", "bullets": []},
            {"slide_index": 1, "type": "closing", "title": "Close", "bullets": []},
        ],
    }
    template = {
        "color_scheme": {},
        "slides": [
            {"index": 0, "layout": "cover", "text_zones": [], "all_zones": []},
            {"index": 1, "layout": "closing", "text_zones": [], "all_zones": []},
        ],
    }
    client = SwappingClient()

    mapping = _llm_mapping(
        client,
        outline,
        {"template_id": "test"},
        {"theme_profile": {}, "slides": []},
        {},
        template,
    )

    assert [slide["template_slide_index"] for slide in mapping["slides"]] == [0, 1]
    assert "content_mapping_retry_0_1" in client.phases
    assert "content_mapping_retry_1_1" in client.phases


def test_micro_repair_context_excludes_project_and_neighbor_payloads() -> None:
    class RecordingClient:
        def __init__(self) -> None:
            self.contexts = []

        def generate_json(self, **kwargs):
            self.contexts.append(kwargs["context"])
            return {"repairs": []}

        def was_fallback(self, _phase):
            return True

    client = RecordingClient()
    slide = {
        "slide_index": 0,
        "template_slide_index": 0,
        "zones": [],
    }
    template_slide = {
        "text_zones": [{
            "zone_id": "title",
            "type": "title",
            "editable": True,
            "original_text": "Original title",
            "min_chars_hint": 4,
            "max_chars_hint": 12,
        }]
    }

    _repair_zones_in_micro_batches(
        client,
        slide,
        {"slide_index": 0, "title": "New title", "bullets": ["Fact one"]},
        [{"slide_index": 1, "title": "large neighbor payload"}],
        {"large": "design payload"},
        {"key_facts": ["large project payload"]},
        template_slide,
    )

    assert client.contexts
    assert all("neighboring_slides" not in context for context in client.contexts)
    assert all("design_decision" not in context for context in client.contexts)
    assert all("neighboring_zone_copy" not in context for context in client.contexts)
    assert all("project_context" not in context for context in client.contexts)


def test_picture_filled_autoshape_is_an_image_zone_and_replaceable(tmp_path) -> None:
    source_image = tmp_path / "source.png"
    replacement_image = tmp_path / "replacement.png"
    Image.new("RGB", (80, 60), "red").save(source_image)
    Image.new("RGB", (80, 60), "blue").save(replacement_image)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    picture = slide.shapes.add_picture(
        str(source_image), Inches(0.2), Inches(0.2), Inches(1), Inches(1)
    )
    filled_shape = slide.shapes.add_shape(
        5, Inches(2), Inches(1), Inches(3), Inches(2)
    )
    # python-pptx lacks a public user-picture fill setter. This mirrors the
    # DrawingML pattern found in authored templates.
    filled_shape._element.spPr.append(deepcopy(picture._element.blipFill))

    zones = _zones_from_pptx_shapes(
        slide, 0, prs.slide_width, prs.slide_height
    )
    zone = next(
        item for item in zones
        if item["native_shape_id"] == filled_shape.shape_id
    )
    assert zone["type"] == "image"
    assert zone["editable"] is False

    blip = filled_shape._element.xpath('.//*[local-name()="blip"]')[0]
    relationship_id = blip.get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    )
    before = filled_shape.part.rels[relationship_id].target_part.blob
    _replace_shape_image(filled_shape, replacement_image)
    after = filled_shape.part.rels[relationship_id].target_part.blob

    assert before != after
    assert after == replacement_image.read_bytes()


def test_cleared_surplus_paragraph_does_not_leave_empty_bullet() -> None:
    from pptx.oxml.xmlchemy import OxmlElement

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    first = shape.text_frame.paragraphs[0]
    first.text = "Old one"
    second = shape.text_frame.add_paragraph()
    second.text = "Old two"
    for paragraph in (first, second):
        p_pr = paragraph._p.get_or_add_pPr()
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "•")
        p_pr.append(bullet)

    _apply_text_to_shape(shape, {"type": "body", "content": "Only item"})

    second_tags = [
        child.tag.rsplit("}", 1)[-1]
        for child in second._p.get_or_add_pPr()
    ]
    assert second.text == ""
    assert "buNone" in second_tags
    assert "buChar" not in second_tags


def test_strict_assembly_honors_selected_template_pages(tmp_path) -> None:
    prs = Presentation()
    template_slides = []
    for index in range(4):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        shape.text = f"Template {index}"
        template_slides.append({
            "index": index,
            "all_zones": [{
                "zone_id": f"slide-{index + 1}/shape-{shape.shape_id}",
                "native_shape_id": shape.shape_id,
                "type": "title",
                "position": [0.1, 0.1, 0.4, 0.1],
            }],
        })
    template_path = tmp_path / "template.pptx"
    output_path = tmp_path / "output.pptx"
    prs.save(template_path)
    slide_contents = {
        "assembly_policy": {"mode": "text_replace_only"},
        "slides": [
            {"slide_index": 0, "template_slide_index": 1, "zones": []},
            {"slide_index": 1, "template_slide_index": 3, "zones": []},
        ],
    }
    mappings = {
        0: {"title": {"zone_id": template_slides[1]["all_zones"][0]["zone_id"], "type": "title", "action": "replace_text", "content": "Output A"}},
        1: {"title": {"zone_id": template_slides[3]["all_zones"][0]["zone_id"], "type": "title", "action": "replace_text", "content": "Output B"}},
    }

    write_pptx_from_mapping(
        template_path,
        {"slides": template_slides},
        mappings,
        {},
        slide_contents,
        output_path,
    )

    output = Presentation(output_path)
    assert len(output.slides) == 2
    assert "Output A" in " ".join(shape.text for shape in output.slides[0].shapes if shape.has_text_frame)
    assert "Output B" in " ".join(shape.text for shape in output.slides[1].shapes if shape.has_text_frame)


def test_text_replace_only_ignores_full_page_background_and_preserves_template(
    tmp_path,
) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    shape.text = "Original"
    zone_id = f"slide-1/shape-{shape.shape_id}"
    template_path = tmp_path / "template.pptx"
    output_path = tmp_path / "output.pptx"
    prs.save(template_path)
    background_dir = tmp_path / "background_images"
    background_dir.mkdir()
    Image.new("RGB", (1600, 900), "red").save(background_dir / "slide-000.png")

    write_pptx_from_mapping(
        template_path,
        {
            "slides": [{
                "index": 0,
                "all_zones": [{
                    "zone_id": zone_id,
                    "native_shape_id": shape.shape_id,
                    "type": "title",
                    "position": [0.1, 0.1, 0.4, 0.1],
                }],
            }]
        },
        {
            0: {
                zone_id: {
                    "zone_id": zone_id,
                    "type": "title",
                    "action": "replace_text",
                    "content": "Replacement",
                }
            }
        },
        {},
        {
            "assembly_policy": {"mode": "text_replace_only"},
            "slides": [{
                "slide_index": 0,
                "template_slide_index": 0,
                "zones": [],
            }],
        },
        output_path,
        workspace_root=tmp_path,
    )

    output = Presentation(output_path)
    output_slide = output.slides[0]
    assert "Replacement" in " ".join(
        item.text for item in output_slide.shapes if item.has_text_frame
    )
    assert all(
        item.shape_type != MSO_SHAPE_TYPE.PICTURE
        for item in output_slide.shapes
    )


def test_generated_background_replace_requires_explicit_assembly_mode(tmp_path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    shape.text = "Template furniture"
    template_path = tmp_path / "template.pptx"
    output_path = tmp_path / "output.pptx"
    prs.save(template_path)
    background_dir = tmp_path / "background_images"
    background_dir.mkdir()
    Image.new("RGB", (1600, 900), "blue").save(background_dir / "slide-000.png")

    write_pptx_from_mapping(
        template_path,
        {"slides": [{"index": 0, "all_zones": []}]},
        {0: {}},
        {},
        {
            "assembly_policy": {"mode": "generated_background_replace"},
            "slides": [{
                "slide_index": 0,
                "template_slide_index": 0,
                "zones": [],
            }],
        },
        output_path,
        workspace_root=tmp_path,
    )

    output = Presentation(output_path)
    output_slide = output.slides[0]
    assert any(
        item.shape_type == MSO_SHAPE_TYPE.PICTURE
        for item in output_slide.shapes
    )
    assert "Template furniture" not in " ".join(
        item.text for item in output_slide.shapes if item.has_text_frame
    )


def test_template_match_requires_an_absolute_score_and_clear_margin() -> None:
    assert not _has_confident_template_match([("health", 0.167), ("travel", 0.165)])
    assert not _has_confident_template_match([("health", 0.7), ("travel", 0.695)])
    assert _has_confident_template_match([("library", 0.7), ("health", 0.5)])


def test_required_template_uses_low_confidence_real_candidate(tmp_path, monkeypatch) -> None:
    template_dir = tmp_path / "templates" / "health"
    template_dir.mkdir(parents=True)
    (template_dir / "template.pptx").write_bytes(b"pptx")
    monkeypatch.chdir(tmp_path)

    selected, score, usable = _choose_required_real_template(
        [("health", 0.167), ("fallback.default", 0.5)],
        [
            {"template_id": "health", "path": "health"},
            {"template_id": "fallback.default", "path": ""},
        ],
    )

    assert selected["template_id"] == "health"
    assert score == 0.167
    assert usable == [("health", 0.167)]


def test_strict_invariants_detect_geometry_drift(tmp_path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    shape.text = "Template heading"
    zones = _zones_from_pptx_shapes(slide, 0, prs.slide_width, prs.slide_height)
    output_path = tmp_path / "output.pptx"
    prs.save(output_path)
    slide_contents = {
        "assembly_policy": {"mode": "text_replace_only"},
        "slides": [{"slide_index": 0, "template_slide_index": 0, "zones": []}],
    }
    template_zones = {"slides": [{"index": 0, "all_zones": zones}]}

    unchanged = _strict_assembly_invariants(output_path, template_zones, slide_contents)
    assert unchanged["applicable"]
    assert unchanged["passed"]

    mutated = Presentation(output_path)
    mutated.slides[0].shapes[0].left += Inches(0.25)
    mutated.save(output_path)

    drifted = _strict_assembly_invariants(output_path, template_zones, slide_contents)
    assert not drifted["passed"]
    assert any(issue["kind"] == "geometry_drift" for issue in drifted["issues"])


def test_content_mapping_capability_exposes_bounded_preview_tool() -> None:
    assert "preview_mapping" in CONTENT_MAPPING.tools


def test_blank_fallback_does_not_render_preserved_image_placeholder(tmp_path) -> None:
    output_path = tmp_path / "fallback.pptx"
    write_pptx({
        "slides": [{
            "slide_index": 0,
            "zones": [
                {"zone_id": "title", "type": "title", "content": "Title", "position": [0.1, 0.1, 0.8, 0.1]},
                {"zone_id": "image", "type": "image", "action": "preserve", "position": [0.7, 0.3, 0.2, 0.4]},
            ],
        }],
    }, output_path)

    presentation = Presentation(output_path)
    all_text = " ".join(
        shape.text for shape in presentation.slides[0].shapes if shape.has_text_frame
    )
    assert "Image placeholder" not in all_text


def test_assembly_rejects_fallback_template_even_with_force(tmp_path) -> None:
    workspace = JobWorkspace(tmp_path)
    workspace.artifact_path("slide_contents").write_text(
        json.dumps({"review_status": "approved", "slides": []}), encoding="utf-8"
    )
    workspace.artifact_path("selected_template").write_text(
        json.dumps({"template_id": "fallback.default"}), encoding="utf-8"
    )

    with pytest.raises(ValueError, match="Real-template assembly is required"):
        run_ppt_assembler(workspace, force=True)
